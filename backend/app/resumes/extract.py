"""Text extraction from uploaded resume files with strict type validation."""

import io

from docx import Document
from pypdf import PdfReader
from pypdf.errors import PdfReadError

ALLOWED = {
    "pdf": {"application/pdf"},
    "docx": {"application/vnd.openxmlformats-officedocument.wordprocessingml.document"},
    "txt": {"text/plain"},
}


class UnsupportedFile(ValueError):
    pass


def detect_type(filename: str, content_type: str | None, data: bytes) -> str:
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext not in ALLOWED:
        raise UnsupportedFile("Only PDF, DOCX or TXT resumes are supported")
    if content_type and content_type not in ALLOWED[ext] and content_type != "application/octet-stream":
        raise UnsupportedFile("File content type does not match its extension")
    # Magic-byte checks so a renamed binary can't slip through.
    if ext == "pdf" and not data.startswith(b"%PDF"):
        raise UnsupportedFile("File is not a valid PDF")
    if ext == "docx" and not data.startswith(b"PK"):
        raise UnsupportedFile("File is not a valid DOCX")
    if ext == "txt" and b"\x00" in data[:4096]:
        raise UnsupportedFile("File is not a plain-text resume")
    return ext


def _decode_text(data: bytes) -> str:
    # UTF-8 first; fall back to Windows-1252, the default "ANSI" encoding of Notepad on Windows.
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("cp1252", errors="replace")


def extract_text(kind: str, data: bytes) -> str:
    try:
        if kind == "pdf":
            reader = PdfReader(io.BytesIO(data))
            return "\n".join((page.extract_text() or "") for page in reader.pages)
        if kind == "docx":
            doc = Document(io.BytesIO(data))
            parts = [p.text for p in doc.paragraphs]
            for table in doc.tables:
                for row in table.rows:
                    parts.append(" | ".join(cell.text for cell in row.cells))
            return "\n".join(parts)
        return _decode_text(data)
    except (PdfReadError, KeyError, ValueError) as exc:
        raise UnsupportedFile("Could not read the file; it may be corrupted") from exc
