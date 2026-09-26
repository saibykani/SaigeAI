"""ATS-friendly DOCX export: single column, standard headings, real text (no tables or images)."""

import io

from docx import Document
from docx.shared import Pt, RGBColor

from app.schemas.resume import ParsedResume

ACCENT = RGBColor(0x1A, 0x73, 0xE8)


def _heading(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text.upper())
    run.bold = True
    run.font.size = Pt(11)
    run.font.color.rgb = ACCENT
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(2)


def _base_doc() -> Document:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)
    for section in doc.sections:
        section.left_margin = section.right_margin = Pt(54)
        section.top_margin = section.bottom_margin = Pt(46)
    return doc


def resume_docx(r: ParsedResume) -> bytes:
    doc = _base_doc()
    name = doc.add_paragraph()
    run = name.add_run(r.name or "Your Name")
    run.bold = True
    run.font.size = Pt(20)
    contact = [x for x in (r.email, r.phone, r.location, r.links.linkedin, r.links.github, r.links.portfolio) if x]
    if contact:
        doc.add_paragraph(" | ".join(contact))
    if r.summary:
        _heading(doc, "Professional Summary")
        doc.add_paragraph(r.summary)
    if r.skills:
        _heading(doc, "Skills")
        doc.add_paragraph(", ".join(r.skills))
    if r.experience:
        _heading(doc, "Experience")
        for e in r.experience:
            p = doc.add_paragraph()
            t = p.add_run(" — ".join(x for x in (e.title, e.company) if x))
            t.bold = True
            dates = " – ".join(x for x in (e.start_date, "Present" if e.is_current else e.end_date) if x)
            meta = " | ".join(x for x in (e.location, dates) if x)
            if meta:
                p.add_run(f"   {meta}")
            for b in e.achievements + e.responsibilities:
                doc.add_paragraph(b, style="List Bullet")
    if r.projects:
        _heading(doc, "Projects")
        for pr in r.projects:
            p = doc.add_paragraph()
            p.add_run(pr.name).bold = True
            if pr.technologies:
                p.add_run(f"   ({', '.join(pr.technologies)})")
            for h in pr.highlights:
                doc.add_paragraph(h, style="List Bullet")
    if r.education:
        _heading(doc, "Education")
        for ed in r.education:
            years = " – ".join(str(y) for y in (ed.start_year, ed.end_year) if y)
            line = ", ".join(x for x in (f"{ed.degree} in {ed.field}" if ed.field else ed.degree, ed.institution, years) if x)
            doc.add_paragraph(line or ed.raw)
    if r.certifications:
        _heading(doc, "Certifications")
        for c in r.certifications:
            doc.add_paragraph(c, style="List Bullet")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def letter_docx(text: str) -> bytes:
    doc = _base_doc()
    for para in text.split("\n\n"):
        doc.add_paragraph(para.strip())
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
